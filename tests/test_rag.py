from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.rag import RAGService, OllamaError, cosine_similarity
from app.main import app
from app.text_processing import TextSegment, chunk_segments, extract_text


class FakeOllama:
    def __init__(self):
        self.generated_prompt = ""

    async def embed(self, text: str) -> list[float]:
        lowered = text.lower()
        return [lowered.count("python"), lowered.count("retrieval"), len(lowered) / 1000]

    async def generate(self, prompt: str) -> str:
        self.generated_prompt = prompt
        return "Grounded answer from the uploaded context."


def test_chunking_keeps_overlap_and_location():
    chunks = chunk_segments([TextSegment("one two three four five six", "Page 2")], chunk_size=4, overlap=1)
    assert [chunk.text for chunk in chunks] == ["one two three four", "four five six"]
    assert all(chunk.location == "Page 2" for chunk in chunks)


def test_cosine_similarity_is_normalized():
    assert cosine_similarity([1, 0], [5, 0]) == pytest.approx(1.0)
    assert cosine_similarity([1, 0], [0, 1]) == pytest.approx(0.0)
    assert cosine_similarity([], [1]) == 0.0


def test_docx_extraction_returns_meaningful_text():
    document = DocxDocument()
    document.add_paragraph("Retrieval augmented generation helps students study.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Loan period"
    table.cell(0, 1).text = "14 days"
    buffer = BytesIO()
    document.save(buffer)
    segments = extract_text("workshop.docx", buffer.getvalue())
    assert len(segments) == 1
    assert "students study" in segments[0].text
    assert "Loan period | 14 days" in segments[0].text
    assert segments[0].location == "Document"


def test_upload_rejects_unsupported_file_type():
    response = TestClient(app).post(
        "/api/documents/upload",
        files={"file": ("notes.txt", b"plain text", "text/plain")},
    )
    assert response.status_code == 415
    assert "PDF or DOCX" in response.json()["detail"]


@pytest.mark.asyncio
async def test_retrieval_ranks_semantically_closest_chunk_and_grounds_prompt():
    client = FakeOllama()
    config = SimpleNamespace(chunk_size=100, chunk_overlap=10, top_k=1)
    service = RAGService(client, config)
    service.documents["doc"] = SimpleNamespace(
        chunks=[
            SimpleNamespace(id="python", document_id="doc", document_name="notes.docx", text="Python retrieval", embedding=[1, 1, 0], location="Document"),
            SimpleNamespace(id="other", document_id="doc", document_name="notes.docx", text="Cooking recipes", embedding=[0, 0, 1], location="Document"),
        ]
    )
    sources = await service.retrieve("python retrieval", top_k=1)
    assert sources[0]["chunk_id"] == "python"
    answer, answer_sources = await service.answer("python retrieval", top_k=1)
    assert answer.startswith("Grounded answer")
    assert answer_sources[0]["chunk_id"] == "python"
    assert "CONTEXT:" in client.generated_prompt
    assert "Cooking recipes" not in client.generated_prompt


def docx_bytes(text):
    document = DocxDocument()
    document.add_paragraph(text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_pdf_extraction_preserves_page_number():
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})
    })
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 10 200 Td (Python retrieval workshop) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    buffer = BytesIO()
    writer.write(buffer)
    segments = extract_text("notes.pdf", buffer.getvalue())
    assert segments == [TextSegment("Python retrieval workshop", "Page 2")]


def test_full_api_upload_retrieve_answer_delete(monkeypatch):
    import app.main as main

    ollama = FakeOllama()
    service = RAGService(ollama)
    monkeypatch.setattr(main, "rag", service)
    client = TestClient(app)
    assert client.get("/").status_code == 200
    assert client.get("/static/app.js").status_code == 200
    assert client.post("/api/chat", json={"question": "python"}).status_code == 422
    response = client.post("/api/documents/upload", files={
        "file": ("lesson.docx", docx_bytes("Python retrieval workshop."))
    })
    assert response.status_code == 200
    document_id = response.json()["document"]["id"]
    assert len(client.get("/api/documents").json()["documents"]) == 1
    retrieved = client.post("/api/retrieve", json={"question": "Python retrieval", "top_k": 1})
    assert retrieved.status_code == 200
    assert retrieved.json()["sources"][0]["text"] == "Python retrieval workshop."
    answer = client.post("/api/chat", json={"question": "Python retrieval"})
    assert answer.status_code == 200
    assert answer.json()["sources"][0]["document_id"] == document_id
    assert "Python retrieval workshop." in ollama.generated_prompt
    assert client.delete(f"/api/documents/{document_id}").status_code == 200
    assert client.get("/api/documents").json() == {"documents": []}
    assert client.post("/api/retrieve", json={"question": "python"}).status_code == 422


@pytest.mark.parametrize("name,data", [("bad.pdf", b"broken"), ("bad.docx", b"broken"), ("empty.docx", b"")])
def test_invalid_upload_has_actionable_error(name, data):
    response = TestClient(app).post("/api/documents/upload", files={"file": (name, data)})
    assert response.status_code in (400, 422)
    assert isinstance(response.json()["detail"], str)


@pytest.mark.asyncio
async def test_embedding_failure_never_adds_partial_document():
    class BrokenOllama(FakeOllama):
        async def embed(self, text):
            raise OllamaError("Model missing")

    service = RAGService(BrokenOllama())
    with pytest.raises(OllamaError):
        await service.add_document("notes.docx", docx_bytes("Python"))
    assert service.documents == {}
